import io
import math
import os
import re
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.db.database import db
from app.db.models import ActivityLog, FinancialRecord, HarvestRecord

IMPORT_ACTION = 'IMPORT_PRODUCTION_EXCEL'
ALLOWED_EXTENSIONS = {'xlsx', 'csv'}
IMPORT_MODES = {'replace', 'append'}
MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_ROWS = 1000
MAX_NOTES_LENGTH = 1000
MAX_AMOUNT = 10 ** 13

PERIOD_PATTERN = re.compile(r'^\d{4}-(0[1-9]|1[0-2])$')

COLUMNS = [
    {'key': 'periode', 'required': True, 'kind': 'period',
     'description': 'Wajib. Format YYYY-MM, contoh 2026-01.'},
    {'key': 'total_produksi_kg', 'required': True, 'kind': 'number',
     'description': 'Wajib. Total produksi dalam kg, angka >= 0 (boleh desimal).'},
    {'key': 'hasil_panen_kg', 'required': False, 'kind': 'number',
     'description': 'Opsional. Hasil panen dalam kg. Jika kosong, mengikuti total_produksi_kg.'},
    {'key': 'biaya_operasional', 'required': False, 'kind': 'number',
     'description': 'Opsional. Biaya operasional dalam Rupiah, tanpa simbol Rp.'},
    {'key': 'pendapatan', 'required': False, 'kind': 'number',
     'description': 'Opsional. Pendapatan dalam Rupiah, tanpa simbol Rp.'},
    {'key': 'catatan', 'required': False, 'kind': 'text',
     'description': 'Opsional. Keterangan tambahan, maksimal 1000 karakter.'},
]
COLUMN_KEYS = [c['key'] for c in COLUMNS]
REQUIRED_KEYS = [c['key'] for c in COLUMNS if c['required']]

TEMPLATE_EXAMPLES = [
    ['2026-01', 12500, 12000, 45000000, 87500000, 'Panen raya blok A'],
    ['2026-02', 9800.5, None, 38000000, 68600000, ''],
]


class ImportFileError(ValueError):
    pass


def build_template(fmt='xlsx'):
    if fmt == 'csv':
        frame = pd.DataFrame(TEMPLATE_EXAMPLES, columns=COLUMN_KEYS, dtype=object)
        buffer = io.BytesIO(frame.to_csv(index=False).encode('utf-8-sig'))
        return buffer, 'text/csv', 'template_impor_produksi.csv'

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Data Produksi'
    sheet.append(COLUMN_KEYS)
    for row in TEMPLATE_EXAMPLES:
        sheet.append(row)

    header_fill = PatternFill('solid', fgColor='1B4332')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    for row in sheet.iter_rows(min_row=2, max_col=1):
        for cell in row:
            cell.number_format = '@'
    for letter, width in zip('ABCDEF', [12, 20, 18, 20, 18, 36]):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = 'A2'

    guide = workbook.create_sheet('Petunjuk')
    guide.append(['Kolom', 'Keterangan'])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for column in COLUMNS:
        guide.append([column['key'], column['description']])
    guide.append([])
    guide.append(['Catatan', 'Tulis angka tanpa pemisah ribuan. Gunakan titik untuk desimal, contoh 9800.5.'])
    guide.append(['Catatan', 'Baris yang seluruh kolomnya kosong akan dilewati.'])
    guide.column_dimensions['A'].width = 20
    guide.column_dimensions['B'].width = 80

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return (
        buffer,
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'template_impor_produksi.xlsx',
    )


def _is_blank(value):
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    if value is pd.NaT:
        return True
    return isinstance(value, str) and not value.strip()


def _display_value(value):
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.strftime('%Y-%m-%d')
    return str(value).strip()


def _normalize_header(value):
    return re.sub(r'[\s\-]+', '_', str(value).strip().lower())


def _parse_period(value):
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return f"{value.year:04d}-{value.month:02d}"
    text = str(value).strip()
    match = re.match(r'^(\d{4})[-/](\d{1,2})(?:[-/]\d{1,2})?(?:[ T].*)?$', text)
    if match:
        text = f"{match.group(1)}-{int(match.group(2)):02d}"
    if not PERIOD_PATTERN.match(text):
        raise ValueError('harus berformat YYYY-MM (contoh 2026-01)')
    return text


def _parse_number(value):
    if isinstance(value, bool):
        raise ValueError('harus berupa angka')
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = re.sub(r'(?i)^rp\.?', '', str(value).strip()).replace(' ', '')
        if ',' in text and '.' in text:
            if text.rfind(',') > text.rfind('.'):
                text = text.replace('.', '').replace(',', '.')
            else:
                text = text.replace(',', '')
        elif text.count(',') > 1:
            text = text.replace(',', '')
        elif ',' in text:
            text = text.replace(',', '.')
        elif text.count('.') > 1:
            text = text.replace('.', '')
        try:
            number = float(text)
        except ValueError:
            raise ValueError('harus berupa angka')
    if math.isnan(number) or math.isinf(number):
        raise ValueError('harus berupa angka')
    if number < 0:
        raise ValueError('tidak boleh negatif')
    if number >= MAX_AMOUNT:
        raise ValueError('nilai terlalu besar')
    return round(number, 2)


def validate_row(raw):
    errors = []
    data = {}
    for column in COLUMNS:
        key = column['key']
        value = raw.get(key)
        if _is_blank(value):
            if column['required']:
                errors.append(f"{key} wajib diisi")
            data[key] = None
            continue
        try:
            if column['kind'] == 'period':
                data[key] = _parse_period(value)
            elif column['kind'] == 'number':
                data[key] = _parse_number(value)
            else:
                text = str(value).strip()
                if len(text) > MAX_NOTES_LENGTH:
                    raise ValueError(f"maksimal {MAX_NOTES_LENGTH} karakter")
                data[key] = text
        except ValueError as e:
            errors.append(f"{key} {e}")
            data[key] = _display_value(value)

    if not errors and data['hasil_panen_kg'] is None:
        data['hasil_panen_kg'] = data['total_produksi_kg']
    return data, errors


def _validate_rows(raw_rows):
    rows = []
    seen_periods = {}
    for row_number, raw in raw_rows:
        data, errors = validate_row(raw)
        period = data.get('periode')
        if period and not errors:
            if period in seen_periods:
                errors.append(f"periode {period} duplikat dengan baris {seen_periods[period]}")
            else:
                seen_periods[period] = row_number
        rows.append({
            'row_number': row_number,
            'status': 'error' if errors else 'valid',
            'errors': errors,
            'data': data,
        })
    return rows


def _summarize(rows, skipped_blank=0):
    valid_rows = [r for r in rows if r['status'] == 'valid']
    return {
        'rows': rows,
        'total_rows': len(rows),
        'valid_count': len(valid_rows),
        'error_count': len(rows) - len(valid_rows),
        'skipped_blank': skipped_blank,
        'total_production_kg': round(sum(r['data']['total_produksi_kg'] for r in valid_rows), 2),
        'errors': [{'row_number': r['row_number'], 'messages': r['errors']} for r in rows if r['errors']],
    }


def _read_frame(file):
    filename = os.path.basename(file.filename or '')
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if ext not in ALLOWED_EXTENSIONS:
        raise ImportFileError('Format file harus .xlsx atau .csv')

    content = file.read(MAX_FILE_SIZE + 1)
    if not content:
        raise ImportFileError('File kosong')
    if len(content) > MAX_FILE_SIZE:
        raise ImportFileError(f"Ukuran file maksimal {MAX_FILE_SIZE // (1024 * 1024)} MB")

    try:
        if ext == 'xlsx':
            return pd.read_excel(io.BytesIO(content), engine='openpyxl', dtype=object, sheet_name=0)
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = content.decode('latin-1')
        return pd.read_csv(io.StringIO(text), dtype=str, keep_default_na=False, sep=None, engine='python')
    except ImportFileError:
        raise
    except Exception:
        raise ImportFileError('File tidak dapat dibaca. Pastikan file tidak rusak dan sesuai template.')


def preview_file(file):
    frame = _read_frame(file)
    frame.columns = [_normalize_header(c) for c in frame.columns]

    missing = [key for key in REQUIRED_KEYS if key not in frame.columns]
    if missing:
        raise ImportFileError(f"Kolom wajib tidak ditemukan: {', '.join(missing)}. Gunakan template yang disediakan.")

    raw_rows = []
    skipped_blank = 0
    for index, record in enumerate(frame.to_dict(orient='records')):
        values = {key: record.get(key) for key in COLUMN_KEYS}
        if all(_is_blank(v) for v in values.values()):
            skipped_blank += 1
            continue
        raw_rows.append((index + 2, values))

    if not raw_rows:
        raise ImportFileError('File tidak berisi baris data')
    if len(raw_rows) > MAX_ROWS:
        raise ImportFileError(f"Maksimal {MAX_ROWS} baris data per impor")

    return _summarize(_validate_rows(raw_rows), skipped_blank)


def commit_import(farm, user, rows, mode='replace', filename=None, ip_address=None):
    if mode not in IMPORT_MODES:
        raise ImportFileError("Mode impor harus 'replace' atau 'append'")
    if not isinstance(rows, list) or not rows:
        raise ImportFileError('Tidak ada baris data untuk disimpan')
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"Maksimal {MAX_ROWS} baris data per impor")
    if not all(isinstance(r, dict) for r in rows):
        raise ImportFileError('Format data baris tidak valid')

    validated = _validate_rows(
        (r['row_number'] if isinstance(r.get('row_number'), int) else i + 1, {key: r.get(key) for key in COLUMN_KEYS})
        for i, r in enumerate(rows)
    )
    summary = _summarize(validated)
    if summary['error_count']:
        return False, summary

    company_id = farm.project.company_id
    periods = sorted({r['data']['periode'] for r in validated})
    safe_filename = os.path.basename(filename or '')[:255] or 'tanpa nama'

    try:
        deleted_financial = deleted_harvest = 0
        if mode == 'replace':
            deleted_financial = FinancialRecord.query.filter(
                FinancialRecord.farm_id == farm.id, FinancialRecord.period.in_(periods)
            ).delete(synchronize_session=False)
            deleted_harvest = HarvestRecord.query.filter(
                HarvestRecord.farm_id == farm.id, HarvestRecord.period.in_(periods)
            ).delete(synchronize_session=False)

        for row in validated:
            data = row['data']
            notes = data['catatan'] or ''
            db.session.add(FinancialRecord(
                company_id=company_id,
                farm_id=farm.id,
                period=data['periode'],
                total_production_kg=data['total_produksi_kg'],
                operational_cost=data['biaya_operasional'] or 0,
                estimated_revenue=data['pendapatan'] or 0,
                notes=notes,
            ))
            db.session.add(HarvestRecord(
                company_id=company_id,
                farm_id=farm.id,
                period=data['periode'],
                yield_kg=data['hasil_panen_kg'],
                notes=notes,
            ))

        mode_label = 'ganti periode' if mode == 'replace' else 'tambah'
        db.session.add(ActivityLog(
            user_id=user.id,
            action=IMPORT_ACTION,
            entity_type='Farm',
            entity_id=farm.id,
            details=(
                f"Mengimpor {len(validated)} catatan produksi & panen dari '{safe_filename}' "
                f"ke lahan '{farm.name}' (mode {mode_label}, periode {', '.join(periods)})"
            ),
            ip_address=ip_address,
        ))
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    return True, {
        'imported': len(validated),
        'periods': periods,
        'mode': mode,
        'deleted_financial': deleted_financial,
        'deleted_harvest': deleted_harvest,
        'total_production_kg': summary['total_production_kg'],
    }
